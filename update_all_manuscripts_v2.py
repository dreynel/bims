import os
import shutil
import docx
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

DIAGRAM_FILES = {
    "Figure 1": {
        "file": os.path.abspath("Documentation/Diagrams/figure1_conceptual_framework.png"),
        "width": 6.2
    },
    "Figure 2": {
        "file": os.path.abspath("Documentation/Diagrams/figure2_prototyping_model.png"),
        "width": 4.0
    },
    "Figure 3": {
        "file": os.path.abspath("Documentation/Diagrams/figure3_use_case_diagram.png"),
        "width": 6.2
    },
    "Figure 4": {
        "file": os.path.abspath("Documentation/Diagrams/figure4_activity_diagram.png"),
        "width": 5.8
    },
    "Figure 5": {
        "file": os.path.abspath("Documentation/Diagrams/figure5_class_diagram.png"),
        "width": 5.8
    },
    "Figure 6": {
        "file": os.path.abspath("Documentation/Diagrams/figure6_sequence_diagram.png"),
        "width": 6.2
    },
    "Figure 7": {
        "file": os.path.abspath("Documentation/Diagrams/figure7_system_architecture.png"),
        "width": 6.2
    }
}

SIGNIFICANCE_HEADINGS = [
    "Punong Barangay (Barangay Captain).",
    "Barangay Secretary (and Clerical Staff).",
    "Residents of Barangay Dacutan.",
    "Municipality of Dumangas and DILG.",
    "Iloilo Science and Technology University - Dumangas Campus (ISAT U - DC).",
    "Iloilo Science and Technology University (ISAT U).",
    "Future Researchers."
]

def set_run_font(run, font_name="Times New Roman", size_pt=12, bold=False, italic=False):
    run.font.name = font_name
    run.font.size = Pt(size_pt)
    run.bold = bold
    run.italic = italic

def replace_image_for_figure(doc, fig_key, image_path, width_in):
    caption_prefix = f"{fig_key}."
    for i, p in enumerate(doc.paragraphs):
        txt = p.text.strip()
        if txt.startswith(caption_prefix):
            # Locate picture paragraph immediately preceding
            prev_pic_p = None
            for j in range(i - 1, max(-1, i - 15), -1):
                if 'drawing' in doc.paragraphs[j]._element.xml:
                    prev_pic_p = doc.paragraphs[j]
                    break
            
            # Create new picture paragraph
            new_p = doc.add_paragraph()
            new_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            new_p.paragraph_format.space_before = Pt(6)
            new_p.paragraph_format.space_after = Pt(6)
            run = new_p.add_run()
            run.add_picture(image_path, width=Inches(width_in))
            
            # Insert before caption paragraph
            p._element.addprevious(new_p._element)
            
            # Remove old picture paragraph if found
            if prev_pic_p is not None:
                parent = prev_pic_p._element.getparent()
                parent.remove(prev_pic_p._element)
                print(f"  [Image] Replaced image for {fig_key}")
            else:
                print(f"  [Image] Inserted new image for {fig_key} (no previous picture found)")
            return

def update_docx_manuscript(doc_path):
    print(f"\n==========================================")
    print(f"Updating DOCX: {doc_path}")
    print(f"==========================================")
    if not os.path.exists(doc_path):
        print(f"File not found: {doc_path}")
        return

    doc = docx.Document(doc_path)

    # 1. Update text paragraphs
    for p in doc.paragraphs:
        txt = p.text.strip()
        if not txt:
            continue

        # --- A. Background of the Study / Public Portal & AI Chatbot ---
        if "To solve these operational challenges, the researchers proposed and developed the Barangay Information and Management" in txt:
            p.text = (
                "To solve these operational challenges, the researchers proposed and developed the Barangay Information and Management "
                "System (BIMS) for Barangay Dacutan, Dumangas, Iloilo. The system is a web-based e-governance platform designed to modernize "
                "and streamline community operations. It supports three distinct user roles: the Resident (general public), the Punong Barangay "
                "(Hon. Stephen James D. Foerster), and the Barangay Secretary (Mrs. Pearl Joy J. Amuan, who also handles administrative records and "
                "clerical operations). The system provides a citizen-facing public web portal where residents can directly apply for official "
                "certificates and clearances across personal, business, and dispute categories, consult an intelligent 24/7 AI Chatbot Assistant "
                "(powered by the Google Gemini API) for procedural advice and requirements, and track request milestones in real time using "
                "an automated tracking reference code. For barangay officials, the system provides a secured administrative portal with modules "
                "for census profiling across the 7 puroks, automated clearance issuance with authentic Philippine Republic standard A4 layouts, "
                "Katarungang Pambarangay blotter dispute mediation tracking, and emergency calamity relief aid distribution monitoring. "
                "To eliminate the administrative overhead of managing citizen user accounts and forgotten passwords while safeguarding the "
                "system against spam, prank submissions, and malicious automated bot entries, the system adopts the recommended compromise of "
                "One-Time Pin (OTP) Verification via Email. When a resident completes the online document appointment form, the system dispatches "
                "a transient 4-digit numeric code to their active email address. Entering this code verifies that the applicant is a genuine "
                "individual with a valid email account, eliminating fake submissions without forcing residents to register and maintain "
                "password-protected accounts."
            )
            for run in p.runs:
                set_run_font(run)
            print("  [Updated] Background of Study with Gemini AI Chatbot")

        # --- B. Objectives of the Study: Include AI Chatbot ---
        elif "provide a public citizen portal where community residents can view barangay information, schedule document appointments" in txt:
            p.text = (
                "\u2022 provide a public citizen portal where community residents can view barangay information, schedule document appointments, "
                "consult an integrated 24/7 AI Chatbot Assistant (utilizing the Google Gemini API) for automated guidance on clearance requirements "
                "and fee schedules under Barangay Ordinance No. 2020-01, apply online for official certificates and documents across personal, "
                "business, and dispute categories, secured through an One-Time Pin (OTP) verification mechanism via Email that transmits a transient "
                "4-digit code to validate active email addresses and eliminate fake/spam entries without requiring password-based account registration, "
                "receive automated email status notifications upon secretary review and captain approval, and track request progress in real time "
                "using a unique reference code;"
            )
            for run in p.runs:
                set_run_font(run)
            print("  [Updated] Objective 3 with Gemini AI Chatbot")

        # --- C. Conceptual Framework Process Stage ---
        elif "The process stage involves the core activities needed to transform the inputs into the desired system outputs" in txt:
            p.text = (
                "The process stage involves the core activities needed to transform the inputs into the desired system outputs through the "
                "Prototyping Model methodology. These include managing administrative user accounts, providing 24/7 conversational constituent "
                "assistance via an AI Chatbot integrated with the Google Gemini API, authenticating citizen appointment requests through 4-digit "
                "OTP verification via Email to prevent spam without password registration, registering and managing resident records across the "
                "seven puroks, processing document requests through a two-tier clerical review by the Barangay Secretary and executive approval "
                "by the Punong Barangay, recording official regulatory fees and statutory exemptions pursuant to Barangay Ordinance No. 2020-01, "
                "managing blotter dispute mediation records, logging calamity relief distribution, dispatching automated SMTP email status updates, "
                "generating authentic Philippine Republic standard A4 clearances, and conducting software quality evaluation under "
                "System Evaluation ( ISO 25010 )."
            )
            for run in p.runs:
                set_run_font(run)
            print("  [Updated] Conceptual Framework Process Stage with AI & System Evaluation ( ISO 25010 )")

        elif "As reflected in the conceptual diagram, the output of the framework is strictly and solely the completed software deliverable" in txt or "The output of the framework is the Barangay Information and Management System" in txt:
            p.text = (
                "As reflected in the conceptual diagram, the output of the framework is strictly and solely the completed software deliverable: "
                "the Barangay Information and Management System (BIMS). The framework illustrates an orthogonal, straight-line architectural flow "
                "where verified inputs pass through prototyping processes to deliver the unified web system, interconnected with a continuous "
                "System Evaluation ( ISO 25010 ) feedback loop. To ensure the system's operational dependability and excellence, the system is "
                "assessed across functional suitability, performance efficiency, compatibility, usability, reliability, security, maintainability, "
                "and portability, with evaluation results continuously channeled back to refine system operations."
            )
            for run in p.runs:
                set_run_font(run)
            print("  [Updated] Conceptual Framework Output Stage")

        # --- D. Significance of the Study (Italic, Unbold) ---
        for heading in SIGNIFICANCE_HEADINGS:
            if txt.startswith(heading):
                body = txt[len(heading):].strip()
                p.text = ""
                r1 = p.add_run(heading + " ")
                set_run_font(r1, bold=False, italic=True)
                r2 = p.add_run(body)
                set_run_font(r2, bold=False, italic=False)
                break

        # --- E. Scope and Limitations: AI Chatbot ---
        if "For community members, a Public Citizen Portal is available, providing secure and convenient access to official barangay" in txt:
            p.text = (
                "For community members, a Public Citizen Portal is available, providing secure and convenient access to official barangay "
                "announcements, leadership directories, an intelligent AI Chatbot Assistant powered by the Google Gemini API for procedural "
                "guidance, and online applications for official certificates and documents across personal, business, and dispute categories "
                "pursuant to Barangay Ordinance No. 2020-01. To prevent spam, prank submissions, and malicious bot entries without imposing the "
                "friction of mandatory account registration, the portal incorporates a lightweight 4-digit One-Time Pin (OTP) verification "
                "mechanism via Email that validates the applicant's active email address prior to official submission. The portal also features "
                "a live four-stage document tracker, allowing citizens to monitor their application progress in real time using a unique tracking "
                "reference code without having to make repeated physical visits to the barangay hall."
            )
            for run in p.runs:
                set_run_font(run)
            print("  [Updated] Scope: Public Citizen Portal with Gemini AI Chatbot")

        # --- F. Project Description: Module 1 ---
        elif txt.startswith("1. Public Citizen Portal & Document Appointment Desk:"):
            p.text = (
                "1. Public Citizen Portal & Document Appointment Desk: A publicly accessible landing page featuring barangay announcements, "
                "leadership directory, emergency hotlines, and an interactive AI Chatbot Assistant (integrated via Google Gemini API) providing 24/7 "
                "constituent guidance on certificate requirements and office procedures. The portal enables online clearance applications for official "
                "document offerings secured by 4-digit One-Time Pin (OTP) verification via Email (validating active email addresses to eliminate spam "
                "without requiring user passwords), accompanied by live 4-stage document tracking (Submitted -> Evaluated -> Approved -> Ready for Release)."
            )
            for run in p.runs:
                set_run_font(run)
            print("  [Updated] Project Description: Module 1 with Gemini AI Chatbot")

        elif "Figure 3 shows the Use Case Diagram of the Barangay Information and Management System" in txt:
            p.text = (
                "Figure 3 shows the Use Case Diagram of the Barangay Information and Management System, illustrating an inline horizontal "
                "architecture with three distinct actors: Resident / Citizen (General Public) on the left, and Mrs. Pearl Joy J. Amuan (Barangay Secretary & "
                "Records Officer) and Hon. Stephen James D. Foerster (Punong Barangay) on the right. Inside the system boundary, the frontline, "
                "administrative, and executive use case packages are structured inline to eliminate empty space and maintain proportionate height, "
                "incorporating the 24/7 AI Chatbot Assistant (Google Gemini API) alongside clearance applications, 4-digit Email OTP verification, "
                "blotter records management, and executive approval authority."
            )
            for run in p.runs:
                set_run_font(run)
            print("  [Updated] Figure 3 Use Case Narrative")

        elif "Figure 7 illustrates the physical and cloud deployment architecture of BIMS" in txt or "Figure 7 illustrates the physical and network deployment architecture of BIMS" in txt:
            p.text = (
                "Figure 7 illustrates the physical and cloud deployment architecture of BIMS across four modular tiers. Public citizens access "
                "the responsive web portal and AI Chatbot Assistant using mobile smartphones or personal computers, while barangay personnel operate "
                "office desktop workstations and network laser printers at the barangay hall. Communications are secured through HTTPS encryption. "
                "The cloud hosting environment runs a Linux web server stack, hosting the core web application engine, automated SMTP Email & OTP "
                "notification services, and an integration gateway with the Google Gemini API for constituent conversational assistance. The backend "
                "interacts securely with a Relational Database Management System (RDBMS) storing master community records."
            )
            for run in p.runs:
                set_run_font(run)
            print("  [Updated] Figure 7 Architecture Narrative")

    # 2. Update Table 1 (Comparison Table) to add AI Chatbot row
    for t in doc.tables:
        if len(t.rows) > 0 and len(t.columns) == 5:
            header_txt = t.rows[0].cells[0].text.strip()
            if "System Features" in header_txt or "Features" in header_txt:
                has_ai_row = any("AI Chatbot" in r.cells[0].text for r in t.rows)
                if not has_ai_row:
                    new_row = t.add_row()
                    vals = ["24/7 AI Chatbot Resident Assistant (Gemini API)", "No", "No", "No", "Yes"]
                    for idx, val in enumerate(vals):
                        new_row.cells[idx].text = val
                        for p in new_row.cells[idx].paragraphs:
                            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if idx > 0 else WD_ALIGN_PARAGRAPH.LEFT
                            for run in p.runs:
                                set_run_font(run, size_pt=10)
                    print("  [Updated Table 1] Added AI Chatbot row")

    # 3. Re-embed all 7 refreshed figures
    for fig_key, info in DIAGRAM_FILES.items():
        replace_image_for_figure(doc, fig_key, info["file"], info["width"])

    # 4. Format all figure captions consistently
    for p in doc.paragraphs:
        txt = p.text.strip()
        for fig_key in DIAGRAM_FILES.keys():
            if txt.startswith(f"{fig_key}."):
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_before = Pt(4)
                p.paragraph_format.space_after = Pt(12)
                parts = txt.split(".", 1)
                if len(parts) == 2:
                    p.text = ""
                    r1 = p.add_run(parts[0] + ". ")
                    set_run_font(r1, bold=False, italic=False)
                    r2 = p.add_run(parts[1].strip())
                    set_run_font(r2, bold=False, italic=True)

    doc.save(doc_path)
    print(f"Saved: {doc_path}")

def update_markdown_manuscript():
    md_path = "Documentation/Barangay_Information_Management_System_Manuscript_Ch1-3.md"
    print(f"\n==========================================")
    print(f"Updating Markdown: {md_path}")
    print(f"==========================================")
    with open(md_path, "r", encoding="utf-8") as f:
        content = f.read()

    # Add Gemini AI Chatbot to Background
    old_bg_portal = "The system provides a citizen-facing public web portal where residents can directly apply for official certificates and clearances across personal, business, and dispute categories, and track request milestones in real time using an automated tracking reference code."
    new_bg_portal = "The system provides a citizen-facing public web portal where residents can directly apply for official certificates and clearances across personal, business, and dispute categories, consult an intelligent 24/7 AI Chatbot Assistant (powered by the Google Gemini API) for procedural advice and clearance requirements, and track request milestones in real time using an automated tracking reference code."
    if old_bg_portal in content:
        content = content.replace(old_bg_portal, new_bg_portal)

    # Add Gemini AI Chatbot to Objective 3
    old_obj3 = "•  provide a public citizen portal where community residents can view barangay information, schedule document appointments, and apply online for official certificates and documents across personal, business, and dispute categories under Barangay Ordinance No. 2020-01, secured through an One-Time Pin (OTP) verification mechanism via Email"
    new_obj3 = "•  provide a public citizen portal where community residents can view barangay information, schedule document appointments, consult an integrated 24/7 AI Chatbot Assistant (utilizing the Google Gemini API) for automated guidance on clearance requirements and fee schedules under Barangay Ordinance No. 2020-01, apply online for official certificates and documents across personal, business, and dispute categories, secured through an One-Time Pin (OTP) verification mechanism via Email"
    if old_obj3 in content:
        content = content.replace(old_obj3, new_obj3)

    # Update Conceptual Framework Process & System Evaluation ( ISO 25010 )
    old_cf_proc = "conducting ISO 25010 software quality evaluation."
    new_cf_proc = "providing 24/7 conversational constituent assistance via an AI Chatbot integrated with the Google Gemini API, and conducting software quality evaluation under System Evaluation ( ISO 25010 )."
    if old_cf_proc in content:
        content = content.replace(old_cf_proc, new_cf_proc)

    old_cf_out = "interconnected with a continuous evaluation feedback loop. To ensure the system's operational dependability and excellence, the system is assessed using the ISO/IEC 25010 software quality model"
    new_cf_out = "interconnected with a continuous System Evaluation ( ISO 25010 ) feedback loop. To ensure the system's operational dependability and excellence, the system is assessed across functional suitability, performance efficiency, compatibility, usability, reliability, security, maintainability, and portability"
    if old_cf_out in content:
        content = content.replace(old_cf_out, new_cf_out)

    # Module 1
    old_mod1 = "1. Public Citizen Portal & Document Appointment Desk: A publicly accessible landing page featuring barangay announcements, leadership directory, and emergency hotlines, with online clearance application for official document offerings secured by 4-digit One-Time Pin (OTP) verification via Email"
    new_mod1 = "1. Public Citizen Portal & Document Appointment Desk: A publicly accessible landing page featuring barangay announcements, leadership directory, emergency hotlines, and an interactive AI Chatbot Assistant (integrated via Google Gemini API) providing 24/7 constituent guidance on certificate requirements and office procedures, with online clearance applications for official document offerings secured by 4-digit One-Time Pin (OTP) verification via Email"
    if old_mod1 in content:
        content = content.replace(old_mod1, new_mod1)

    # Table 1: add AI Chatbot row
    if "24/7 AI Chatbot Resident Assistant" not in content:
        old_table_end = "| Defined Roles: Resident, Captain, Secretary/Clerk | Yes | Yes | Yes | Yes |"
        new_table_end = "| Defined Roles: Resident, Captain, Secretary/Clerk | Yes | Yes | Yes | Yes |\n| 24/7 AI Chatbot Resident Assistant (Gemini API) | No | No | No | Yes |"
        content = content.replace(old_table_end, new_table_end)

    with open(md_path, "w", encoding="utf-8") as f:
        f.write(content)
    print(f"Saved: {md_path}")

if __name__ == "__main__":
    docx_files = [
        "Documentation/BIMS - updated (latest).docx",
        "Documentation/Barangay_Information_Management_System_Manuscript_Ch1-3.docx",
        "Documentation/BIMS - updated.docx"
    ]
    for d in docx_files:
        update_docx_manuscript(d)
    update_markdown_manuscript()
    print("\nALL DOCUMENTS AND DIAGRAMS SUCCESSFULLY SYNCHRONIZED AND UP TO DATE!")

