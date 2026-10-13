import os
import shutil
import re
import docx
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

DIAGRAM_FILES = {
    "Figure 1": {
        "file": os.path.abspath("Documentation/Diagrams/figure1_conceptual_framework.png"),
        "width": 6.0
    },
    "Figure 2": {
        "file": os.path.abspath("Documentation/Diagrams/figure2_prototyping_model.png"),
        "width": 3.6
    },
    "Figure 3": {
        "file": os.path.abspath("Documentation/Diagrams/figure3_use_case_diagram.png"),
        "width": 5.8
    },
    "Figure 4": {
        "file": os.path.abspath("Documentation/Diagrams/figure4_activity_diagram.png"),
        "width": 5.8
    },
    "Figure 5": {
        "file": os.path.abspath("Documentation/Diagrams/figure5_class_diagram.png"),
        "width": 5.8
    },
    "Figure 6": {
        "file": os.path.abspath("Documentation/Diagrams/figure6_sequence_diagram.png"),
        "width": 6.0
    },
    "Figure 7": {
        "file": os.path.abspath("Documentation/Diagrams/figure7_system_architecture.png"),
        "width": 6.0
    }
}

def set_run_font(run, font_name="Times New Roman", size_pt=12, bold=False, italic=False):
    run.font.name = font_name
    run.font.size = Pt(size_pt)
    run.bold = bold
    run.italic = italic

def replace_image_for_figure(doc, fig_key, image_path, width_in):
    caption_prefix = f"{fig_key}."
    for i, p in enumerate(doc.paragraphs):
        txt = p.text.strip()
        if txt.startswith(caption_prefix):
            # Locate picture paragraph immediately preceding
            prev_pic_p = None
            for j in range(i - 1, max(-1, i - 15), -1):
                if 'drawing' in doc.paragraphs[j]._element.xml:
                    prev_pic_p = doc.paragraphs[j]
                    break
            
            # Create new picture paragraph
            new_p = doc.add_paragraph()
            new_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            new_p.paragraph_format.space_before = Pt(6)
            new_p.paragraph_format.space_after = Pt(6)
            run = new_p.add_run()
            run.add_picture(image_path, width=Inches(width_in))
            
            # Insert before caption paragraph
            p._element.addprevious(new_p._element)
            
            # Remove old picture paragraph if found
            if prev_pic_p is not None:
                parent = prev_pic_p._element.getparent()
                parent.remove(prev_pic_p._element)
                print(f"  [Image] Replaced image for {fig_key}")
            else:
                print(f"  [Image] Inserted new image for {fig_key} (no previous picture found)")
            return

def update_docx_manuscript(doc_path):
    print(f"\n==========================================")
    print(f"Processing: {doc_path}")
    print(f"==========================================")
    
    if not os.path.exists(doc_path):
        print(f"File not found: {doc_path}")
        return

    backup_path = doc_path.replace(".docx", "_backup_official.docx")
    shutil.copyfile(doc_path, backup_path)
    print(f"Backup created at: {backup_path}")

    doc = docx.Document(doc_path)

    # 1. Update Paragraphs
    for p in doc.paragraphs:
        txt = p.text
        if not txt.strip():
            continue

        orig_txt = txt

        # --- A. Acknowledgements & Officials ---
        if "Hon. Roberto S. Alen" in txt or "Mrs. Ma. Elena Balleza" in txt:
            txt = txt.replace("Hon. Roberto S. Alen", "Hon. Stephen James D. Foerster")
            txt = txt.replace("Roberto S. Alen", "Hon. Stephen James D. Foerster")
            txt = txt.replace("Mrs. Ma. Elena Balleza", "Mrs. Pearl Joy J. Amuan")
            txt = txt.replace("Ma. Elena Balleza", "Mrs. Pearl Joy J. Amuan")

        if "Roberto Dacutan" in txt:
            txt = txt.replace("Hon. Roberto Dacutan", "Hon. Stephen James D. Foerster")
            txt = txt.replace("Roberto Dacutan", "Hon. Stephen James D. Foerster")

        # --- B. Background of Study / Geography / 2025 Demographics ---
        if "Barangay Dacutan is one of the coastal and agricultural barangays situated in the Municipality of Dumangas" in txt and "Purok 1 to Purok 7" in txt:
            txt = (
                "Barangay Dacutan is one of the coastal and agricultural barangays situated in the Municipality of Dumangas, "
                "Province of Iloilo, under Region VI (Western Visayas). Geographically located approximately 2.5 kilometers from the "
                "municipal poblacion with a total land area of 230.7171 hectares, the barangay is bounded on the North by Barangay Sulangan "
                "and Barangay Binaobao, on the South by the coastal fishing waters, on the East by Barangay Talusan and Barangay Sapao, "
                "and on the West by Barangay Ermita. Administratively, the community is organized strictly into seven (7) puroks, designated "
                "as Purok 1 through Purok 7. According to the official 2025 Barangay Demographic Profile, Barangay Dacutan has an aggregate "
                "population of 2,337 residents (comprising 1,142 males and 1,195 females), distributed across 678 families and 562 households. "
                "The community includes significant special sectors: 556 senior citizens (aged 60 and above), 69 persons with disabilities (PWDs), "
                "1,445 registered COMELEC voters, 69 Pantawid Pamilyang Pilipino Program (4Ps) beneficiaries, 75 overseas Filipino workers (OFWs), "
                "52 registered farmers, and 146 active fisherfolks. Frontline regulatory and administrative governance is administered "
                "pursuant to Republic Act No. 7160 (Local Government Code of 1991) and Barangay Ordinance No. 2020-01 ('An Ordinance Enacting the "
                "Barangay Tax and Regulatory Fees of Barangay Dacutan, Dumangas, Iloilo'), enacted by the Sangguniang Barangay under the leadership "
                "of Punong Barangay Hon. Stephen James D. Foerster and Barangay Secretary Mrs. Pearl Joy J. Amuan."
            )

        # --- C. Background of Study / OTP Compromise (Email Only) ---
        if "To solve these problems, the researchers proposed and developed the Barangay Information and Management System (BIMS)" in txt:
            txt = (
                "To solve these operational challenges, the researchers proposed and developed the Barangay Information and Management "
                "System (BIMS) for Barangay Dacutan, Dumangas, Iloilo. The system is a web-based e-governance platform designed to modernize "
                "and streamline community operations. It supports three distinct user roles: the Resident (general public), the Punong Barangay "
                "(Hon. Stephen James D. Foerster), and the Barangay Secretary (Mrs. Pearl Joy J. Amuan, who also handles administrative records and "
                "clerical operations). The system provides a citizen-facing public web portal where residents can directly apply for official "
                "certificates and clearances across personal, business, and dispute categories, and track request milestones in real time using "
                "an automated tracking reference code. For barangay officials, the system provides a secured administrative portal with modules "
                "for census profiling across the 7 puroks, automated clearance issuance with authentic Philippine Republic standard A4 layouts, "
                "Katarungang Pambarangay blotter dispute mediation tracking, and emergency calamity relief aid distribution monitoring. "
                "To eliminate the administrative overhead of managing citizen user accounts and forgotten passwords while safeguarding the "
                "system against spam, prank submissions, and malicious automated bot entries, the system adopts the recommended compromise of "
                "One-Time Pin (OTP) Verification via Email. When a resident completes the online document appointment form, the system dispatches "
                "a transient 4-digit numeric code to their active email address. Entering this code verifies that the applicant is a genuine "
                "individual with a valid email account, eliminating fake submissions without forcing residents to register and maintain "
                "password-protected accounts."
            )

        # --- D. Objectives of the Study ---
        if "manage user accounts and enforce a two-tier review and approval workflow" in txt or "manage secure administrative user accounts for authorized barangay officials" in txt:
            txt = (
                "\u2022 manage secure administrative user accounts for authorized barangay officials (Punong Barangay and Barangay Secretary) "
                "and enforce a two-tier review and approval workflow where citizen document requests are evaluated by the Barangay Secretary "
                "(acting as administrative records officer) and granted final executive approval by the Punong Barangay;"
            )

        if "provide a public citizen portal where community residents can view barangay information" in txt:
            txt = (
                "\u2022 provide a public citizen portal where community residents can view barangay information, schedule document appointments, "
                "and apply online for official certificates and documents across personal, business, and dispute categories under Barangay Ordinance "
                "No. 2020-01, secured through an One-Time Pin (OTP) verification mechanism via Email that transmits a transient 4-digit code to "
                "validate active email addresses and eliminate fake/spam entries without requiring password-based account registration, receive "
                "automated email status notifications upon secretary review and captain approval, and track request progress in real time using "
                "a unique reference code;"
            )

        if "record and monitor resident census records with complete demographic details" in txt:
            txt = (
                "\u2022 record and monitor resident census records with complete demographic details across the seven (7) puroks of Barangay "
                "Dacutan, and track special sector eligibility flags including Senior Citizens, PWDs, 4Ps members, Solo Parents, Farmers, "
                "Fisherfolks, and registered COMELEC voters;"
            )

        # --- E. Conceptual Framework IPO Narrative ---
        if "The input stage consists of the essential data and information required by the system" in txt:
            txt = (
                "The input stage consists of the essential data and information required by the system, including administrative leadership "
                "profiles of Punong Barangay Hon. Stephen James D. Foerster and Barangay Secretary Mrs. Pearl Joy J. Amuan, official resident "
                "census records across the seven (7) puroks of Barangay Dacutan (totaling 2,337 residents, 678 families, and 562 households with "
                "demographic sector flags for senior citizens, PWDs, 4Ps, farmers, and fisherfolks), citizen online document request and appointment "
                "information under Barangay Ordinance No. 2020-01 (including active applicant email addresses and transient 4-digit Email OTP "
                "verification codes), Lupon Tagapamayapa dispute blotter records, and emergency calamity aid distribution logs. These inputs serve "
                "as the foundation for the system's operations and ensure that all processes are based on complete and accurate community data."
            )

        if "The process stage involves the core activities needed to transform the inputs into the desired system outputs" in txt:
            txt = (
                "The process stage involves the core activities needed to transform the inputs into the desired system outputs through the "
                "Prototyping Model methodology. These include managing administrative user accounts, authenticating citizen appointment "
                "requests through 4-digit OTP verification via Email to prevent spam without password registration, registering and managing "
                "resident records across the seven puroks, processing document requests through a two-tier clerical review by the Barangay "
                "Secretary and executive approval by the Punong Barangay, recording official regulatory fees and statutory exemptions pursuant "
                "to Barangay Ordinance No. 2020-01, managing blotter dispute mediation records, logging calamity relief distribution, dispatching "
                "automated SMTP email status updates, generating authentic Philippine Republic standard A4 clearances, and conducting ISO 25010 "
                "software quality evaluation."
            )

        if "The output of the framework is the Barangay Information and Management System" in txt or "output of the framework is strictly and solely" in txt:
            txt = (
                "As reflected in the conceptual diagram, the output of the framework is strictly and solely the completed software deliverable: "
                "the Barangay Information and Management System (BIMS). The framework illustrates an orthogonal, straight-line architectural "
                "flow where verified inputs pass through prototyping processes to deliver the unified web system, interconnected with a "
                "continuous evaluation feedback loop. To ensure the system's operational dependability and excellence, the system is assessed "
                "using the ISO/IEC 25010 software quality model across functional suitability, performance efficiency, compatibility, usability, "
                "reliability, security, maintainability, and portability, with evaluation feedback continuously channeled back to refine system operations."
            )

        # --- F. Definition of Terms ---
        if txt.startswith("One-Time Pin (OTP) Verification."):
            txt = (
                "One-Time Pin (OTP) Verification. A temporary security and authentication mechanism where a transient 4-digit numeric code "
                "is generated and transmitted via email to verify that an applicant possesses an active, valid email address before an online "
                "transaction is submitted (Brillantes & Perante-Calina, 2018). In this study, OTP verification refers to the automated spam prevention "
                "and identity verification feature embedded into the public document appointment form that validates the requestor as a real "
                "person and eliminates fake or bot entries without requiring residents to register and manage password-protected user accounts."
            )

        if txt.startswith("Purok."):
            txt = (
                "Purok. A sub-division or neighborhood cluster within a Philippine barangay that groups residential households for administrative "
                "and census purposes (Brillantes & Perante-Calina, 2018). In this study, purok refers to the seven (7) territorial groupings in "
                "Barangay Dacutan, specifically Purok 1 through Purok 7, used for filtering resident records, mapping demographics, and generating "
                "community analytics."
            )

        # --- G. Significance of the Study ---
        if txt.startswith("Punong Barangay (Barangay Captain)."):
            txt = (
                "Punong Barangay (Barangay Captain). The system provides Punong Barangay Hon. Stephen James D. Foerster with an executive "
                "dashboard showing real-time demographic statistics across the 7 puroks, monthly revenue collections under Ordinance No. 2020-01, "
                "pending dispute rates, and active citizen clearance requests. This provides executive visibility, accelerates Tier-2 clearance "
                "approvals, and supports accurate data-driven community planning."
            )

        if txt.startswith("Barangay Secretary (and Clerical Staff)."):
            txt = (
                "Barangay Secretary (and Clerical Staff). The system significantly lessens the daily administrative and clerical workload of "
                "Barangay Secretary Mrs. Pearl Joy J. Amuan. Instead of manually searching through physical logbooks and paper cards, the secretary "
                "can verify resident records in seconds, process online clearance requests across official document offerings, record official "
                "receipts and CTC numbers, and monitor mediation schedules with automated tracking numbers. In addition, because the system employs "
                "lightweight 4-digit OTP verification via email on appointment requests rather than mandatory citizen account registration, "
                "the secretary is completely relieved from the administrative burden of managing citizen login accounts, handling forgotten "
                "password requests, and cleaning up spam submissions."
            )

        if txt.startswith("Residents of Barangay Dacutan."):
            txt = (
                "Residents of Barangay Dacutan. Community constituents save substantial time and travel expenses by accessing the public portal "
                "directly from their home computers or smartphones. They can request official clearances, certifications, and appointments and "
                "monitor application status in real time without making repeated trips to the barangay hall. Furthermore, residents enjoy a "
                "frictionless requesting experience without having to create or remember account passwords, verifying their identity easily "
                "through a 4-digit OTP sent to their verified email inbox."
            )

        # --- H. Scope and Limitations ---
        if txt.startswith("This study focuses on the development of the Barangay Information and Management System (BIMS) for Barangay Dacutan"):
            txt = (
                "This study focuses on the development of the Barangay Information and Management System (BIMS) for Barangay Dacutan, "
                "Municipality of Dumangas, Province of Iloilo, designed to improve the efficiency, transparency, and delivery of community "
                "records management and frontline public services. The system is intended for use within the administrative jurisdiction of "
                "Barangay Dacutan, specifically covering all residential households situated across Purok 1 through Purok 7."
            )

        if txt.startswith("For community members, a Public Citizen Portal is available"):
            txt = (
                "For community members, a Public Citizen Portal is available, providing secure and convenient access to official barangay "
                "announcements, leadership directories, and online applications for official certificates and documents across personal, "
                "business, and dispute categories pursuant to Barangay Ordinance No. 2020-01. To prevent spam, prank submissions, and malicious "
                "bot entries without imposing the friction of mandatory account registration, the portal incorporates a lightweight 4-digit "
                "One-Time Pin (OTP) verification mechanism via Email that validates the applicant's active email address prior to official "
                "submission. The portal also features a live four-stage document tracker, allowing citizens to monitor their application "
                "progress in real time using a unique tracking reference code without having to make repeated physical visits to the barangay hall."
            )

        if "All administrative users are required to log in with their assigned role-based credentials" in txt or "To ensure accessibility and usability, the system is deployed on a live web hosting" in txt:
            txt = (
                "To ensure accessibility and usability, the system is deployed on a live web hosting environment and accessed online via "
                "the internet through standard web browsers on smartphones, tablets, laptops, and desktop computers. This enables community "
                "residents to access the public citizen portal and submit document requests online 24/7 from anywhere without having to travel "
                "to the barangay hall during the initial application stage. To protect database integrity while keeping citizen services frictionless, "
                "password-based user accounts are strictly restricted to authorized administrative personnel (Punong Barangay Hon. Stephen James D. Foerster "
                "and Barangay Secretary Mrs. Pearl Joy J. Amuan), whereas citizen requestors authenticate their appointments on demand via 4-digit Email OTP."
            )

        if "The system is limited to web-based access and does not include a native downloadable mobile application" in txt or "Cellular SMS messaging is specifically focused" in txt or "The system does not process cellular SMS" in txt:
            txt = (
                "The system is limited to web-based access and does not include a native downloadable mobile application. Users access the system "
                "through modern desktop or mobile web browsers. To prevent ongoing telecommunication overhead and recurring subscriber costs, "
                "the system does not utilize cellular SMS text messaging; instead, all identity verification operations (4-digit Email OTP), "
                "milestone status tracking, and clearance pickup notifications are dispatched through automated SMTP email notifications and "
                "the public online tracking portal. The system does not include an online electronic payment gateway, such as GCash or Maya, "
                "so all document fees are collected personally in cash at the barangay hall upon claiming, recorded with manual Official Receipt (OR) "
                "and Community Tax Certificate (CTC) numbers pursuant to Barangay Ordinance No. 2020-01. Furthermore, the application is strictly "
                "designed for the administrative jurisdiction and demographic records of Barangay Dacutan, Dumangas, Iloilo, and does not "
                "automatically integrate with external municipal or national police law enforcement databases."
            )

        # --- I. Chapter III Project Description & Roles ---
        if "Resident (General Public): Community constituents who access the public portal" in txt:
            txt = (
                "\u2022 Resident (General Public): Community constituents who access the public portal to read announcements, submit appointments "
                "and applications for official document offerings authenticated via a 4-digit OTP sent to their email inbox without needing password "
                "registration, and track processing status online."
            )

        if "Punong Barangay (Barangay Captain): The executive leader who monitors demographic indicators" in txt:
            txt = (
                "\u2022 Punong Barangay (Hon. Stephen James D. Foerster): The executive leader who monitors demographic indicators, approves "
                "certificate releases (Tier-2 approval), oversees Lupon disputes, and authorizes official certifications."
            )

        if "Barangay Secretary (Administrative & Records Officer): The operational administrator who also executes" in txt:
            txt = (
                "\u2022 Barangay Secretary (Mrs. Pearl Joy J. Amuan): The operational administrator and records officer who also executes general "
                "clerical routines, including census resident encoding across the 7 puroks, certificate evaluation (Tier-1 review), fee/OR logging "
                "under Ordinance No. 2020-01, blotter mediation tracking, and calamity aid registry."
            )

        if txt.startswith("1. Public Citizen Portal & Document Appointment Desk:"):
            txt = (
                "1. Public Citizen Portal & Document Appointment Desk: A publicly accessible landing page featuring barangay announcements, "
                "leadership directory, and emergency hotlines, with online clearance application for official document offerings secured by "
                "4-digit One-Time Pin (OTP) verification via Email (validating active email addresses to eliminate spam without requiring user passwords), "
                "accompanied by live 4-stage document tracking (Submitted -> Evaluated -> Approved -> Ready for Release)."
            )

        if "Figure 3 shows the Use Case Diagram of the Barangay Information and Management System" in txt:
            txt = (
                "Figure 3 shows the Use Case Diagram of the Barangay Information and Management System, illustrating the three distinct "
                "actors: Resident / Citizen (General Public), Mrs. Pearl Joy J. Amuan (Barangay Secretary & Records Officer), and Hon. Stephen James D. Foerster "
                "(Punong Barangay), along with the specific system actions and frontline e-services each role performs."
            )

        if "Figure 4 shows the Activity Diagram illustrating the sequential flow of an online document request" in txt:
            txt = (
                "Figure 4 shows the Activity Diagram illustrating the sequential flow of an online document request and appointment, "
                "from initial citizen submission and 4-digit Email OTP verification to clerical review by Barangay Secretary Mrs. Pearl Joy J. Amuan, "
                "executive approval by Punong Barangay Hon. Stephen James D. Foerster, OR/CTC recording under Ordinance No. 2020-01, and final clearance printing."
            )

        if "Figure 7 illustrates the physical and network deployment architecture of BIMS" in txt or "Figure 7 illustrates the physical and cloud deployment architecture" in txt:
            txt = (
                "Figure 7 illustrates the physical and cloud deployment architecture of BIMS, showing how user devices connect over the "
                "internet to the live cloud web server and relational database. Public citizens access the responsive web portal using mobile "
                "smartphones, laptops, or personal computers, while barangay personnel operate office desktop workstations and network laser "
                "printers at the barangay hall. All internet communications are protected through secure HTTPS encryption (Port 443). The cloud "
                "server runs on a Linux (Ubuntu) environment with Nginx/Apache and PHP 8.2, hosting the Laravel 12 application core and an "
                "Automated Email & OTP Dispatcher Service via SMTP. The application layer handles request routing, business logic, role-based "
                "permissions, and database operations with the MySQL 8.0 relational database."
            )

        if "During the Planning phase, the researchers established the project goals" in txt:
            txt = (
                "During the Planning phase, the researchers established the project goals, determined technical and operational feasibility, "
                "and conducted initial coordination with the officials of Barangay Dacutan. The team met with Punong Barangay Hon. Stephen James D. Foerster "
                "and Barangay Secretary Mrs. Pearl Joy J. Amuan to review Barangay Ordinance No. 2020-01, analyze current challenges in managing resident "
                "records and issuing certificates, and evaluate the 2025 demographic census summary. This phase defined the scope, resource requirements, "
                "and schedule for developing the web-based management system."
            )

        if "2. Domain End-Users (5 respondents):" in txt:
            txt = (
                "2. Domain End-Users (5 respondents): Punong Barangay Hon. Stephen James D. Foerster, Barangay Secretary Mrs. Pearl Joy J. Amuan, "
                "and community residents from Barangay Dacutan who evaluated functional suitability, user interface ease of use, clearance generation "
                "accuracy across official document offerings, and practical community applicability."
            )

        # Clean any remaining "SMS" or "mobile phone or email" or "Sitio Ilawod"
        txt = txt.replace("SMS or Email", "Email")
        txt = txt.replace("SMS or email", "email")
        txt = txt.replace("SMS/Email", "Email")
        txt = txt.replace("via SMS", "via Email")
        txt = txt.replace("mobile phone or email", "email address")
        txt = txt.replace("Sitio Ilawod, and Sitio Ilaya", "")
        txt = txt.replace("Sitio Ilawod and Sitio Ilaya", "")
        txt = txt.replace("Purok 1 through Purok 7, Sitio Ilawod, and Sitio Ilaya", "Purok 1 through Purok 7")
        txt = txt.replace("Purok 1 to Purok 7, Sitio Ilawod, and Sitio Ilaya", "Purok 1 to Purok 7")

        if txt != orig_txt:
            # Check if paragraph has bold prefix like "Term. Definition"
            if "." in orig_txt[:40] and any(orig_txt.startswith(k) for k in [
                "Punong Barangay", "Barangay Secretary", "Residents of Barangay Dacutan",
                "One-Time Pin", "Purok."
            ]):
                parts = txt.split(".", 1)
                p.text = ""
                r1 = p.add_run(parts[0] + ".")
                set_run_font(r1, bold=True)
                r2 = p.add_run(parts[1])
                set_run_font(r2, bold=False)
            else:
                p.text = txt
                for run in p.runs:
                    set_run_font(run)
            print(f"  [Updated P] {txt[:70]}...")

    # 2. Update Table Cells (e.g. Appendix C transmittal box)
    for t_idx, table in enumerate(doc.tables):
        for r_idx, row in enumerate(table.rows):
            for c_idx, cell in enumerate(row.cells):
                cell_txt = cell.text
                if not cell_txt.strip():
                    continue
                orig_cell_txt = cell_txt
                cell_txt = cell_txt.replace("Hon. Roberto S. Alen", "Hon. Stephen James D. Foerster")
                cell_txt = cell_txt.replace("Roberto S. Alen", "Hon. Stephen James D. Foerster")
                cell_txt = cell_txt.replace("Mrs. Ma. Elena Balleza", "Mrs. Pearl Joy J. Amuan")
                cell_txt = cell_txt.replace("Ma. Elena Balleza", "Mrs. Pearl Joy J. Amuan")
                cell_txt = cell_txt.replace("SMS or Email", "Email")
                cell_txt = cell_txt.replace("SMS/Email", "Email")
                cell_txt = cell_txt.replace("via SMS", "via Email")
                if cell_txt != orig_cell_txt:
                    cell.text = cell_txt
                    print(f"  [Updated Table {t_idx} R{r_idx}C{c_idx}] {cell_txt[:70]}...")

    # 3. Refresh All 7 Embedded Diagram Images
    for fig_key, info in DIAGRAM_FILES.items():
        replace_image_for_figure(doc, fig_key, info["file"], info["width"])

    # 4. Format all figure captions consistently
    for p in doc.paragraphs:
        txt = p.text.strip()
        for fig_key in DIAGRAM_FILES.keys():
            if txt.startswith(f"{fig_key}."):
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_before = Pt(4)
                p.paragraph_format.space_after = Pt(12)
                parts = txt.split(".", 1)
                if len(parts) == 2:
                    p.text = ""
                    r1 = p.add_run(parts[0] + ". ")
                    set_run_font(r1, bold=False, italic=False)
                    r2 = p.add_run(parts[1].strip())
                    set_run_font(r2, bold=False, italic=True)

    doc.save(doc_path)
    print(f"Successfully saved updated docx: {doc_path}")

if __name__ == "__main__":
    targets = [
        "Documentation/BIMS - updated (latest).docx",
        "Documentation/Barangay_Information_Management_System_Manuscript_Ch1-3.docx",
        "Documentation/BIMS - updated.docx"
    ]
    for target in targets:
        update_docx_manuscript(target)
    print("\nALL 3 CAPSTONE MANUSCRIPT DOCX FILES SUCCESSFULLY SYNCHRONIZED!")

