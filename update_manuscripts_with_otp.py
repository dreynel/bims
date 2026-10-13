import os
import shutil
import docx
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

DIAGRAM_FILES = {
    "Figure 1": {
        "file": os.path.abspath("Documentation/Diagrams/figure1_conceptual_framework.png"),
        "width": 6.0
    },
    "Figure 2": {
        "file": os.path.abspath("Documentation/Diagrams/figure2_prototyping_model.png"),
        "width": 3.6
    },
    "Figure 3": {
        "file": os.path.abspath("Documentation/Diagrams/figure3_use_case_diagram.png"),
        "width": 5.8
    },
    "Figure 4": {
        "file": os.path.abspath("Documentation/Diagrams/figure4_activity_diagram.png"),
        "width": 5.8
    },
    "Figure 5": {
        "file": os.path.abspath("Documentation/Diagrams/figure5_class_diagram.png"),
        "width": 5.8
    },
    "Figure 6": {
        "file": os.path.abspath("Documentation/Diagrams/figure6_sequence_diagram.png"),
        "width": 6.0
    },
    "Figure 7": {
        "file": os.path.abspath("Documentation/Diagrams/figure7_system_architecture.png"),
        "width": 6.0
    }
}

def set_run_font(run, font_name="Times New Roman", size_pt=12, bold=False, italic=False):
    run.font.name = font_name
    run.font.size = Pt(size_pt)
    run.bold = bold
    run.italic = italic

def replace_image_for_figure(doc, fig_key, image_path, width_in):
    caption_prefix = f"{fig_key}."
    for i, p in enumerate(doc.paragraphs):
        txt = p.text.strip()
        if txt.startswith(caption_prefix):
            # Locate picture paragraph immediately preceding
            prev_pic_p = None
            for j in range(i - 1, max(-1, i - 15), -1):
                if 'drawing' in doc.paragraphs[j]._element.xml:
                    prev_pic_p = doc.paragraphs[j]
                    break
            
            # Create new picture paragraph
            new_p = doc.add_paragraph()
            new_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            new_p.paragraph_format.space_before = Pt(6)
            new_p.paragraph_format.space_after = Pt(6)
            run = new_p.add_run()
            run.add_picture(image_path, width=Inches(width_in))
            
            # Insert before caption paragraph
            p._element.addprevious(new_p._element)
            
            # Remove old picture paragraph if found
            if prev_pic_p is not None:
                parent = prev_pic_p._element.getparent()
                parent.remove(prev_pic_p._element)
                print(f"  [Image] Replaced image for {fig_key}")
            else:
                print(f"  [Image] Inserted new image for {fig_key} (no previous picture found)")
            return

def update_document(doc_path):
    print(f"\n==========================================")
    print(f"Updating Document: {doc_path}")
    print(f"==========================================")
    
    if not os.path.exists(doc_path):
        print(f"File not found: {doc_path}")
        return

    backup_path = doc_path.replace(".docx", "_backup_otp.docx")
    shutil.copyfile(doc_path, backup_path)
    print(f"Backup created at: {backup_path}")

    doc = docx.Document(doc_path)

    # 1. Background of the Study: Add OTP Compromise
    bg_updated = False
    for p in doc.paragraphs:
        txt = p.text.strip()
        if "track the real-time status of their requests using an automated reference number" in txt:
            if "One-Time Pin (OTP) Verification" not in txt:
                otp_bg_addition = (
                    " To eliminate the administrative headache of managing citizen user accounts "
                    "and forgotten passwords while safeguarding the system against spam and malicious "
                    "bot entries, the system adopts the recommended compromise of One-Time Pin (OTP) "
                    "Verification via SMS or Email. When a resident completes the online document request "
                    "and appointment form, the system dispatches a transient 4-digit code to their active mobile phone "
                    "or email. Entering this code verifies that the requestor is a genuine individual with an active "
                    "contact number, eliminating fake submissions without forcing residents to register and manage "
                    "password-protected accounts."
                )
                p.text = txt + otp_bg_addition
                for run in p.runs:
                    set_run_font(run)
                bg_updated = True
                print("  [Updated] Background of the study with OTP compromise")
                break

    # 2. Objectives of the Study: Objective 2 & 3
    for p in doc.paragraphs:
        txt = p.text.strip()
        if "manage user accounts and enforce a two-tier review and approval workflow" in txt:
            p.text = (
                "\u2022 manage secure administrative user accounts for authorized barangay officials "
                "and enforce a two-tier review and approval workflow where citizen document requests are "
                "evaluated by the Barangay Secretary (acting as administrative records officer) and "
                "granted final executive approval by the Punong Barangay;"
            )
            for run in p.runs:
                set_run_font(run)
            print("  [Updated] Objective 2 (Administrative accounts)")
        
        elif "provide a public citizen portal where community residents can view barangay information, apply online for twelve" in txt:
            p.text = (
                "\u2022 provide a public citizen portal where community residents can view barangay information, "
                "schedule document appointments, and apply online for twelve (12) official certificates and documents "
                "across personal, business, and dispute categories, secured through an One-Time Pin (OTP) verification "
                "mechanism via SMS or Email that transmits a 4-digit code to validate active contact numbers and eliminate "
                "fake/spam entries without requiring password-based account registration, receive automated email status "
                "notifications upon secretary review and captain approval, and track request progress in real time using "
                "a unique reference code;"
            )
            for run in p.runs:
                set_run_font(run)
            print("  [Updated] Objective 3 (Public citizen portal with OTP)")

    # 3. Conceptual Framework: Input and Process
    for p in doc.paragraphs:
        txt = p.text.strip()
        if "The input stage consists of the essential data and information required by the system, including user information" in txt:
            p.text = (
                "The input stage consists of the essential data and information required by the system, "
                "including administrative user information, resident census information, document request "
                "and appointment information (including citizen contact numbers and 4-digit OTP verification codes), "
                "blotter incident information, and calamity aid information. These inputs serve as the foundation "
                "for the system's operations and ensure that all processes are based on complete and accurate community data."
            )
            for run in p.runs:
                set_run_font(run)
            print("  [Updated] Conceptual Framework (Input stage)")
        
        elif "The process stage involves the core activities needed to transform the inputs into the desired system outputs. These include managing user accounts" in txt:
            p.text = (
                "The process stage involves the core activities needed to transform the inputs into the desired "
                "system outputs. These include managing administrative user accounts, authenticating citizen appointment "
                "requests through 4-digit OTP verification via SMS or email to prevent spam without password registration, "
                "registering and managing resident records, processing document requests through a two-tier clerical review "
                "and executive approval hierarchy, managing blotter and dispute records, recording calamity aid distribution, "
                "and generating system reports. These processes are designed to ensure that the system operates efficiently, "
                "provides reliable records, and speeds up daily administrative tasks at the barangay hall."
            )
            for run in p.runs:
                set_run_font(run)
            print("  [Updated] Conceptual Framework (Process stage)")

    # 4. Definition of Terms: Add One-Time Pin (OTP) Verification
    has_otp_def = any("One-Time Pin (OTP) Verification" in p.text for p in doc.paragraphs)
    if not has_otp_def:
        # Find where to insert (after Laravel)
        for i, p in enumerate(doc.paragraphs):
            txt = p.text.strip()
            if txt.startswith("Laravel."):
                otp_p = doc.add_paragraph()
                p._element.addnext(otp_p._element)
                
                # Format with bold term prefix
                r1 = otp_p.add_run("One-Time Pin (OTP) Verification. ")
                set_run_font(r1, bold=True)
                
                r2 = otp_p.add_run(
                    "A temporary security and authentication mechanism where a transient 4-digit numeric code "
                    "is generated and transmitted via SMS or email to verify that an applicant possesses an active, "
                    "valid contact number before an online transaction is submitted (Brillantes & Perante-Calina, 2018). "
                    "In this study, OTP verification refers to the automated spam prevention and identity verification "
                    "feature embedded into the public document appointment form that validates the requestor as a real "
                    "person and eliminates fake or bot entries without requiring residents to register and manage "
                    "password-protected user accounts."
                )
                set_run_font(r2, bold=False)
                print("  [Inserted] Definition of Terms: One-Time Pin (OTP) Verification")
                break

    # 5. Significance of the Study: Secretary and Residents
    for p in doc.paragraphs:
        txt = p.text.strip()
        if txt.startswith("Barangay Secretary (and Clerical Staff)."):
            if "lightweight 4-digit OTP verification" not in txt:
                p.text = (
                    "Barangay Secretary (and Clerical Staff). The system significantly lessens the daily administrative "
                    "and clerical workload of the secretary. Instead of manually searching through physical logbooks and "
                    "index cards, the secretary can verify resident records in seconds, process online clearance requests "
                    "across 12 document categories, record official receipts, and monitor mediation schedules with automated "
                    "tracking numbers. In addition, because the system employs lightweight 4-digit OTP verification on appointment "
                    "requests rather than mandatory citizen account registration, the secretary is completely relieved from the "
                    "administrative burden of managing citizen login accounts, handling forgotten password requests, and cleaning "
                    "up spam submissions."
                )
                parts = p.text.split(".", 1)
                p.text = ""
                r1 = p.add_run(parts[0] + ".")
                set_run_font(r1, bold=True)
                r2 = p.add_run(parts[1])
                set_run_font(r2, bold=False)
                print("  [Updated] Significance: Barangay Secretary")

        elif txt.startswith("Residents of Barangay Dacutan."):
            if "frictionless requesting experience" not in txt:
                p.text = (
                    "Residents of Barangay Dacutan. Community members save substantial time and travel expenses by "
                    "accesscribing the public portal directly from their home computers or smartphones. They can request any of the 12 "
                    "offered personal, business, or dispute documents and monitor application status in real time without making repeated "
                    "trips to the barangay hall. Furthermore, residents enjoy a frictionless requesting experience without having to "
                    "create or remember account passwords, verifying their identity easily through a 4-digit OTP sent to their mobile "
                    "phone or email."
                )
                parts = p.text.split(".", 1)
                p.text = ""
                r1 = p.add_run(parts[0] + ".")
                set_run_font(r1, bold=True)
                r2 = p.add_run(parts[1])
                set_run_font(r2, bold=False)
                print("  [Updated] Significance: Residents of Barangay Dacutan")

    # 6. Scope and Limitations of the Study
    for p in doc.paragraphs:
        txt = p.text.strip()
        if txt.startswith("For community members, a Public Citizen Portal is available"):
            if "lightweight 4-digit One-Time Pin (OTP)" not in txt:
                p.text = (
                    "For community members, a Public Citizen Portal is available, providing secure and convenient access to "
                    "official barangay announcements, leadership directories, and online applications for twelve (12) official certificates "
                    "and documents across personal, business, and dispute categories. To prevent spam, prank submissions, and malicious "
                    "bot entries without imposing the friction of mandatory account registration, the portal incorporates a lightweight "
                    "4-digit One-Time Pin (OTP) verification mechanism via SMS or Email that validates the applicant's active contact number "
                    "prior to official submission. The portal also features a live four-stage document tracker, allowing citizens to monitor "
                    "their application progress in real time using a unique tracking reference code without having to make repeated physical "
                    "visits to the barangay hall."
                )
                for run in p.runs:
                    set_run_font(run)
                print("  [Updated] Scope: Public Citizen Portal with OTP")

        elif "All administrative users are required to log in with their assigned role-based credentials" in txt:
            p.text = (
                "To ensure accessibility and usability, the system is deployed on a live web hosting environment and accessed "
                "online via the internet through standard web browsers on smartphones, tablets, laptops, and desktop computers. "
                "This enables community residents to access the public citizen portal and submit document requests online 24/7 "
                "from anywhere without having to travel to the barangay hall during the initial application stage. To protect database "
                "integrity while keeping citizen services frictionless, password-based user accounts are strictly restricted to "
                "authorized administrative personnel (Punong Barangay and Barangay Secretary), whereas citizen requestors authenticate "
                "their appointments on demand via 4-digit OTP."
            )
            for run in p.runs:
                set_run_font(run)
            print("  [Updated] Scope: Administrative credentials vs Citizen OTP")

        elif "The system does not process cellular SMS text alerts due to recurring subscription fees" in txt:
            p.text = (
                "The system is limited to web-based access and does not include a native downloadable mobile application. "
                "Users must access the system through modern desktop or mobile web browsers. Cellular SMS messaging is specifically "
                "focused on sending short, transient 4-digit OTP verification codes to validate applicant contact numbers upon submission, "
                "while general status milestone updates and clearance pickup notices are dispatched via automated SMTP email notifications "
                "and the public online tracking portal to minimize ongoing SMS gateway overhead. The system does not include an online "
                "electronic payment gateway, such as GCash or Maya, so all document fees are collected personally in cash at the barangay "
                "hall upon claiming, recorded with manual Official Receipt and Community Tax Certificate numbers. Furthermore, the "
                "application is strictly designed for the administrative jurisdiction and demographic records of Barangay Dacutan, "
                "Dumangas, Iloilo, and does not automatically integrate with external municipal or national police law enforcement databases."
            )
            for run in p.runs:
                set_run_font(run)
            print("  [Updated] Limitations: SMS OTP scope and Email tracking")

    # 7. Chapter III: Project Description
    for p in doc.paragraphs:
        txt = p.text.strip()
        if "Resident (General Public): Community constituents who access the public portal" in txt:
            p.text = (
                "\u2022 Resident (General Public): Community constituents who access the public portal to read announcements, "
                "submit appointments and applications for 12 document offerings authenticated via a 4-digit OTP sent to their mobile "
                "phone or email without needing password registration, and track processing status online."
            )
            for run in p.runs:
                set_run_font(run)
            print("  [Updated] Project Description: Resident Role")

        elif txt.startswith("1. Public Citizen Portal:") or txt.startswith("1. Public Citizen Portal & Document Appointment Desk:"):
            p.text = (
                "1. Public Citizen Portal & Document Appointment Desk: A publicly accessible landing page featuring barangay announcements, "
                "leadership directory, and emergency hotlines, with online clearance application for 12 document offerings secured by 4-digit "
                "One-Time Pin (OTP) verification via SMS or Email (validating active contact numbers to eliminate spam without requiring user passwords), "
                "accompanied by live 4-stage document tracking (Submitted -> Evaluated -> Approved -> Ready for Release)."
            )
            for run in p.runs:
                set_run_font(run)
            print("  [Updated] Project Description: Module 1")

        elif "Figure 4 shows the Activity Diagram illustrating the sequential flow of an online document request" in txt:
            p.text = (
                "Figure 4 shows the Activity Diagram illustrating the sequential flow of an online document request and appointment, "
                "from initial citizen submission and 4-digit OTP verification to secretary evaluation, captain approval, and final printing."
            )
            for run in p.runs:
                set_run_font(run)
            print("  [Updated] Activity Diagram narrative")

    # 8. Refresh Embedded Images for Figures 1, 3, 4, and 7
    replace_image_for_figure(doc, "Figure 1", DIAGRAM_FILES["Figure 1"]["file"], DIAGRAM_FILES["Figure 1"]["width"])
    replace_image_for_figure(doc, "Figure 3", DIAGRAM_FILES["Figure 3"]["file"], DIAGRAM_FILES["Figure 3"]["width"])
    replace_image_for_figure(doc, "Figure 4", DIAGRAM_FILES["Figure 4"]["file"], DIAGRAM_FILES["Figure 4"]["width"])
    replace_image_for_figure(doc, "Figure 7", DIAGRAM_FILES["Figure 7"]["file"], DIAGRAM_FILES["Figure 7"]["width"])

    # 9. Format all figure captions consistently
    for p in doc.paragraphs:
        txt = p.text.strip()
        for fig_key in DIAGRAM_FILES.keys():
            if txt.startswith(f"{fig_key}."):
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_before = Pt(4)
                p.paragraph_format.space_after = Pt(12)
                parts = txt.split(".", 1)
                if len(parts) == 2:
                    p.text = ""
                    r1 = p.add_run(parts[0] + ". ")
                    set_run_font(r1, bold=False, italic=False)
                    r2 = p.add_run(parts[1].strip())
                    set_run_font(r2, bold=False, italic=True)

    doc.save(doc_path)
    print(f"Successfully saved updated docx: {doc_path}")

if __name__ == "__main__":
    update_document("Documentation/BIMS - updated (latest).docx")
    update_document("Documentation/Barangay_Information_Management_System_Manuscript_Ch1-3.docx")
    update_document("Documentation/BIMS - updated.docx")
    print("\nALL CAPSTONE MANUSCRIPT DOCUMENTS SUCCESSFULLY SYNCHRONIZED WITH OTP COMPROMISE!")

