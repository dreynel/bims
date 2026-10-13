import os
import shutil
import docx
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

DIAGRAM_FILES = {
    "Figure 1": {
        "file": os.path.abspath("Documentation/Diagrams/figure1_conceptual_framework.png"),
        "width": 6.2
    },
    "Figure 2": {
        "file": os.path.abspath("Documentation/Diagrams/figure2_prototyping_model.png"),
        "width": 3.8
    },
    "Figure 3": {
        "file": os.path.abspath("Documentation/Diagrams/figure3_use_case_diagram.png"),
        "width": 5.8
    },
    "Figure 4": {
        "file": os.path.abspath("Documentation/Diagrams/figure4_activity_diagram.png"),
        "width": 5.8
    },
    "Figure 5": {
        "file": os.path.abspath("Documentation/Diagrams/figure5_class_diagram.png"),
        "width": 5.8
    },
    "Figure 6": {
        "file": os.path.abspath("Documentation/Diagrams/figure6_sequence_diagram.png"),
        "width": 6.2
    },
    "Figure 7": {
        "file": os.path.abspath("Documentation/Diagrams/figure7_system_architecture.png"),
        "width": 6.2
    }
}

SIGNIFICANCE_HEADINGS = [
    "Punong Barangay (Barangay Captain).",
    "Barangay Secretary (and Clerical Staff).",
    "Residents of Barangay Dacutan.",
    "Municipality of Dumangas and DILG.",
    "Iloilo Science and Technology University - Dumangas Campus (ISAT U - DC).",
    "Iloilo Science and Technology University (ISAT U).",
    "Future Researchers."
]

def set_run_font(run, font_name="Times New Roman", size_pt=12, bold=False, italic=False):
    run.font.name = font_name
    run.font.size = Pt(size_pt)
    run.bold = bold
    run.italic = italic

def replace_image_for_figure(doc, fig_key, image_path, width_in):
    caption_prefix = f"{fig_key}."
    for i, p in enumerate(doc.paragraphs):
        txt = p.text.strip()
        if txt.startswith(caption_prefix):
            # Locate picture paragraph immediately preceding
            prev_pic_p = None
            for j in range(i - 1, max(-1, i - 15), -1):
                if 'drawing' in doc.paragraphs[j]._element.xml:
                    prev_pic_p = doc.paragraphs[j]
                    break
            
            # Create new picture paragraph
            new_p = doc.add_paragraph()
            new_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            new_p.paragraph_format.space_before = Pt(6)
            new_p.paragraph_format.space_after = Pt(6)
            run = new_p.add_run()
            run.add_picture(image_path, width=Inches(width_in))
            
            # Insert before caption paragraph
            p._element.addprevious(new_p._element)
            
            # Remove old picture paragraph if found
            if prev_pic_p is not None:
                parent = prev_pic_p._element.getparent()
                parent.remove(prev_pic_p._element)
                print(f"  [Image] Replaced image for {fig_key}")
            else:
                print(f"  [Image] Inserted new image for {fig_key} (no previous picture found)")
            return

def update_docx(doc_path):
    print(f"\n==========================================")
    print(f"Updating: {doc_path}")
    print(f"==========================================")
    if not os.path.exists(doc_path):
        print(f"File not found: {doc_path}")
        return

    doc = docx.Document(doc_path)

    # 1. Update Significance of the Study headings to unbold, italic
    in_significance = False
    for p in doc.paragraphs:
        txt = p.text.strip()
        if txt.startswith("Significance of the Study"):
            in_significance = True
            continue
        if in_significance and txt.startswith("Scope and Limitations"):
            in_significance = False
            continue

        if in_significance:
            for heading in SIGNIFICANCE_HEADINGS:
                if txt.startswith(heading):
                    body = txt[len(heading):].strip()
                    p.text = ""
                    r1 = p.add_run(heading + " ")
                    set_run_font(r1, bold=False, italic=True)
                    r2 = p.add_run(body)
                    set_run_font(r2, bold=False, italic=False)
                    print(f"  [Significance] Formatted as italic, unbold: {heading}")
                    break

    # 2. Re-embed all 7 updated figures
    for fig_key, info in DIAGRAM_FILES.items():
        replace_image_for_figure(doc, fig_key, info["file"], info["width"])

    # 3. Format all figure captions consistently
    for p in doc.paragraphs:
        txt = p.text.strip()
        for fig_key in DIAGRAM_FILES.keys():
            if txt.startswith(f"{fig_key}."):
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_before = Pt(4)
                p.paragraph_format.space_after = Pt(12)
                parts = txt.split(".", 1)
                if len(parts) == 2:
                    p.text = ""
                    r1 = p.add_run(parts[0] + ". ")
                    set_run_font(r1, bold=False, italic=False)
                    r2 = p.add_run(parts[1].strip())
                    set_run_font(r2, bold=False, italic=True)

    doc.save(doc_path)
    print(f"Saved: {doc_path}")

def update_markdown():
    md_path = "Documentation/Barangay_Information_Management_System_Manuscript_Ch1-3.md"
    print(f"\nUpdating Markdown: {md_path}")
    with open(md_path, "r", encoding="utf-8") as f:
        content = f.read()

    for heading in SIGNIFICANCE_HEADINGS:
        # replace "Heading Body" with "*Heading* Body" if not already italicized
        plain_pattern = f"\n{heading}"
        italic_replacement = f"\n*{heading}*"
        content = content.replace(plain_pattern, italic_replacement)
        content = content.replace(f"\n**{heading}**", italic_replacement)

    with open(md_path, "w", encoding="utf-8") as f:
        f.write(content)
    print(f"Saved: {md_path}")

if __name__ == "__main__":
    docx_files = [
        "Documentation/BIMS - updated (latest).docx",
        "Documentation/Barangay_Information_Management_System_Manuscript_Ch1-3.docx",
        "Documentation/BIMS - updated.docx"
    ]
    for d in docx_files:
        update_docx(d)
    update_markdown()
    print("\nALL DOCUMENTS AND DIAGRAMS FULLY UPDATED AND SYNCHRONIZED!")
