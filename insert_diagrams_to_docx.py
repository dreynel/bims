import os
import shutil
import docx
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

DIAGRAMS = {
    "Figure 1": {
        "file": os.path.abspath("Documentation/Diagrams/figure1_conceptual_framework.png"),
        "width": 6.0
    },
    "Figure 2": {
        "file": os.path.abspath("Documentation/Diagrams/figure2_prototyping_model.png"),
        "width": 3.6
    },
    "Figure 3": {
        "file": os.path.abspath("Documentation/Diagrams/figure3_use_case_diagram.png"),
        "width": 5.8
    },
    "Figure 4": {
        "file": os.path.abspath("Documentation/Diagrams/figure4_activity_diagram.png"),
        "width": 5.8
    },
    "Figure 5": {
        "file": os.path.abspath("Documentation/Diagrams/figure5_class_diagram.png"),
        "width": 5.8
    },
    "Figure 6": {
        "file": os.path.abspath("Documentation/Diagrams/figure6_sequence_diagram.png"),
        "width": 6.0
    },
    "Figure 7": {
        "file": os.path.abspath("Documentation/Diagrams/figure7_system_architecture.png"),
        "width": 6.0
    }
}

def process_manuscript(doc_path):
    print(f"\n==========================================")
    print(f"Processing Manuscript: {doc_path}")
    print(f"==========================================")
    backup_path = doc_path.replace(".docx", "_backup.docx")
    shutil.copyfile(doc_path, backup_path)
    print(f"Backup created at: {backup_path}")

    doc = docx.Document(doc_path)

    # Tables to replace:
    # Scan all tables for Figure placeholders
    tables_to_replace = []
    for t_idx, table in enumerate(doc.tables):
        first_cell_txt = table.rows[0].cells[0].text.strip()
        for fig_key, fig_info in DIAGRAMS.items():
            if f"[{fig_key}." in first_cell_txt:
                tables_to_replace.append((t_idx, table, fig_key, fig_info))
                break

    print(f"Found {len(tables_to_replace)} figure placeholder tables to replace.")

    for t_idx, table, fig_key, fig_info in tables_to_replace:
        print(f"Replacing Table {t_idx} with {fig_key} ({fig_info['file']}, width={fig_info['width']} in)...")
        # Create a new paragraph
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(6)
        run = p.add_run()
        run.add_picture(fig_info["file"], width=Inches(fig_info["width"]))
        
        # Move paragraph before table
        table._element.addprevious(p._element)
        
        # Remove table from document
        parent = table._element.getparent()
        parent.remove(table._element)

    # Ensure all Figure captions are centered and properly italicized
    for p in doc.paragraphs:
        txt = p.text.strip()
        for fig_key in DIAGRAMS.keys():
            if txt.startswith(f"{fig_key}."):
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                # Ensure spacing
                p.paragraph_format.space_before = Pt(4)
                p.paragraph_format.space_after = Pt(12)
                # Ensure the title is italic non-bold
                parts = txt.split(".", 1)
                if len(parts) == 2:
                    p.text = ""
                    r1 = p.add_run(parts[0] + ". ")
                    r1.bold = False
                    r1.italic = False
                    r1.font.name = "Times New Roman"
                    r1.font.size = Pt(12)
                    r2 = p.add_run(parts[1].strip())
                    r2.bold = False
                    r2.italic = True
                    r2.font.name = "Times New Roman"
                    r2.font.size = Pt(12)

    doc.save(doc_path)
    print(f"Successfully saved updated manuscript: {doc_path}")


def process_bims_updated(doc_path):
    print(f"\n==========================================")
    print(f"Processing BIMS Updated: {doc_path}")
    print(f"==========================================")
    backup_path = doc_path.replace(".docx", "_backup.docx")
    shutil.copyfile(doc_path, backup_path)
    print(f"Backup created at: {backup_path}")

    doc = docx.Document(doc_path)

    # In BIMS - updated.docx:
    # 1. Figure 1 does NOT have a table, but caption "Figure 1. Conceptual Framework of the Study (IPO Model)"
    # Find caption for Figure 1 and insert picture right before it.
    for i, p in enumerate(doc.paragraphs):
        txt = p.text.strip()
        if txt.startswith("Figure 1."):
            print(f"Found Figure 1 caption at paragraph {i}. Inserting Figure 1 image before it...")
            img_p = doc.add_paragraph()
            img_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            img_p.paragraph_format.space_before = Pt(6)
            img_p.paragraph_format.space_after = Pt(6)
            run = img_p.add_run()
            run.add_picture(DIAGRAMS["Figure 1"]["file"], width=Inches(DIAGRAMS["Figure 1"]["width"]))
            p._element.addprevious(img_p._element)
            
            # Clean up empty paragraphs immediately preceding the image
            # Look backwards from p in the parent XML
            prev = img_p._element.getprevious()
            while prev is not None and prev.tag.endswith('p') and (prev.text is None or prev.text.strip() == ''):
                to_remove = prev
                prev = prev.getprevious()
                p._element.getparent().remove(to_remove)
            break

    # 2. Figures 2 to 7 have placeholder tables:
    tables_to_replace = []
    for t_idx, table in enumerate(doc.tables):
        first_cell_txt = table.rows[0].cells[0].text.strip()
        for fig_key, fig_info in DIAGRAMS.items():
            if fig_key == "Figure 1":
                continue
            if f"[{fig_key}." in first_cell_txt:
                tables_to_replace.append((t_idx, table, fig_key, fig_info))
                break

    print(f"Found {len(tables_to_replace)} figure placeholder tables in BIMS - updated to replace.")

    for t_idx, table, fig_key, fig_info in tables_to_replace:
        print(f"Replacing Table {t_idx} with {fig_key} ({fig_info['file']}, width={fig_info['width']} in)...")
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(6)
        p.paragraph_format.space_after = Pt(6)
        run = p.add_run()
        run.add_picture(fig_info["file"], width=Inches(fig_info["width"]))
        
        table._element.addprevious(p._element)
        parent = table._element.getparent()
        parent.remove(table._element)

    # Format captions
    for p in doc.paragraphs:
        txt = p.text.strip()
        for fig_key in DIAGRAMS.keys():
            if txt.startswith(f"{fig_key}."):
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_before = Pt(4)
                p.paragraph_format.space_after = Pt(12)
                parts = txt.split(".", 1)
                if len(parts) == 2:
                    p.text = ""
                    r1 = p.add_run(parts[0] + ". ")
                    r1.bold = False
                    r1.italic = False
                    r1.font.name = "Times New Roman"
                    r1.font.size = Pt(12)
                    r2 = p.add_run(parts[1].strip())
                    r2.bold = False
                    r2.italic = True
                    r2.font.name = "Times New Roman"
                    r2.font.size = Pt(12)

    doc.save(doc_path)
    print(f"Successfully saved updated BIMS doc: {doc_path}")

if __name__ == "__main__":
    process_bims_updated("Documentation/BIMS - updated.docx")
    process_manuscript("Documentation/Barangay_Information_Management_System_Manuscript_Ch1-3.docx")
    print("\nALL DOCUMENTS UPDATED WITH OFFICIAL B&W DIAGRAMS SUCCESSFULLY!")
